"""Build the MediTwin project deck (presentation/MediTwin.pptx).

Run:  uv run --with python-pptx python presentation/make_deck.py
Screenshots in presentation/img/ come from the Streamlit command centre.
"""

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "img"
FONT = "Calibri"

BG = RGBColor(0x07, 0x0B, 0x16)
PANEL = RGBColor(0x10, 0x17, 0x2A)
LINE = RGBColor(0x24, 0x30, 0x4D)
TEXT = RGBColor(0xE6, 0xEB, 0xF5)
MUTED = RGBColor(0x94, 0xA0, 0xB8)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
PINK = RGBColor(0xF4, 0x72, 0xB6)
VIOLET = RGBColor(0xA7, 0x8B, 0xFA)
AMBER = RGBColor(0xFB, 0xBF, 0x24)
GREEN = RGBColor(0x34, 0xD3, 0x99)
RED = RGBColor(0xF8, 0x71, 0x71)
GREY = RGBColor(0x64, 0x74, 0x8B)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
TOTAL = 16
_n = 0


# ------------------------------------------------------------------ helpers
def text(slide, x, y, w, h, runs, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, overrides) tuples."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for seg in ([(para, {})] if isinstance(para, str) else para):
            t, o = seg if isinstance(seg, tuple) else (seg, {})
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=PANEL, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, lw=1):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def picture(slide, path, x, y, w=None, h=None, border=LINE):
    pic = slide.shapes.add_picture(str(path), x, y, w, h)
    pic.line.color.rgb = border
    pic.line.width = Pt(1)
    return pic


def new_slide(title=None, kicker=None, accent=CYAN):
    global _n
    _n += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    # soft glow in the corner
    g = s.shapes.add_shape(MSO_SHAPE.OVAL, W - Inches(4.2), -Inches(2.6), Inches(6.5), Inches(5))
    g.fill.solid()
    g.fill.fore_color.rgb = accent
    _alpha(g, 7)
    g.line.fill.background()
    g.shadow.inherit = False
    if title:
        if kicker:
            text(s, Inches(0.6), Inches(0.42), Inches(10), Inches(0.3), kicker.upper(), size=12, color=accent,
                 bold=True)
        text(s, Inches(0.6), Inches(0.7), Inches(12), Inches(0.7), title, size=32, bold=True)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.42), Inches(0.9), Emu(38000))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()
    if _n > 1:
        text(s, Inches(0.6), H - Inches(0.45), Inches(6), Inches(0.3), "MediTwin · AI-Powered Hospital Digital Twin",
             size=10, color=GREY)
        text(s, W - Inches(1.6), H - Inches(0.45), Inches(1), Inches(0.3), f"{_n:02d} / {TOTAL}", size=10,
             color=GREY, align=PP_ALIGN.RIGHT)
    return s


def _alpha(shape, pct):
    """Set fill transparency (pct = opacity %)."""
    clr = shape.fill._xPr.find(qn("a:solidFill"))[0]
    a = clr.makeelement(qn("a:alpha"), {"val": str(pct * 1000)})
    clr.append(a)


def bullets(slide, x, y, w, items, size=16, color=TEXT, bullet_color=CYAN, gap=8):
    tb = slide.shapes.add_textbox(x, y, w, Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.1
        r = p.add_run()
        r.text = "▸  "
        r.font.name, r.font.size, r.font.color.rgb, r.font.bold = FONT, Pt(size), bullet_color, True
        segs = it if isinstance(it, list) else [it]
        for seg in segs:
            t, o = seg if isinstance(seg, tuple) else (seg, {})
            r = p.add_run()
            r.text = t
            r.font.name, r.font.size = FONT, Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.color.rgb = o.get("color", color)
    return tb


def stat(slide, x, y, w, h, value, label, color=CYAN, sub=None):
    box(slide, x, y, w, h)
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Inches(0.18), y, w - Inches(0.36), Emu(30000))
    strip.fill.solid()
    strip.fill.fore_color.rgb = color
    strip.line.fill.background()
    text(slide, x + Inches(0.2), y + Inches(0.18), w - Inches(0.4), Inches(0.3), label.upper(), size=10,
         color=MUTED, bold=True)
    text(slide, x + Inches(0.2), y + Inches(0.45), w - Inches(0.4), Inches(0.7), value, size=30, color=color,
         bold=True)
    if sub:
        text(slide, x + Inches(0.2), y + h - Inches(0.42), w - Inches(0.4), Inches(0.3), sub, size=11, color=MUTED)


def pill(slide, x, y, label, color, w=None):
    w = w or Inches(0.12 * len(label) + 0.4)
    p = box(slide, x, y, w, Inches(0.34), fill=PANEL, line=color, radius=0.5)
    tf = p.text_frame
    tf.margin_top = tf.margin_bottom = 0
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(11), True, color
    return w


def node(slide, x, y, w, h, title, sub, color):
    b = box(slide, x, y, w, h, fill=PANEL, line=color, lw=1.75)
    tf = b.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(14), True, color
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r = p2.add_run()
    r.text = sub
    r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(10.5), MUTED
    return b


def arrow(slide, x1, y1, x2, y2, color=MUTED):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def table(slide, x, y, w, rows, col_w, header_color=CYAN, size=12, row_h=0.36, highlight_col=None, left_cols=(0,), colors=None):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(row_h * len(rows)))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    for k in ("firstRow", "bandRow"):
        tblPr.set(k, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No style, table grid"
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0x16, 0x20, 0x3A) if i == 0 else (PANEL if i % 2 else BG)
            c.margin_left = c.margin_right = Inches(0.1)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                tcPr = c._tc.get_or_add_tcPr()
                ln = tcPr.makeelement(qn(edge), {"w": "6350"})
                sf = ln.makeelement(qn("a:solidFill"), {})
                sf.append(sf.makeelement(qn("a:srgbClr"), {"val": "24304D"}))
                ln.append(sf)
                tcPr.append(ln)
            tf = c.text_frame
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j in left_cols else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            r.font.name, r.font.size = FONT, Pt(size)
            r.font.bold = i == 0 or j == highlight_col
            r.font.color.rgb = header_color if i == 0 else (GREEN if j == highlight_col else TEXT)
            if i and colors and j in colors and val in colors[j]:
                r.font.color.rgb, r.font.bold = colors[j][val], True
    return tbl


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


# ------------------------------------------------------------------ 1 title
s = new_slide()
s.shapes.add_picture(str(IMG / "home.png"), Inches(6.4), Inches(0.9), height=Inches(5.7))
veil = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.4), Inches(0.9), Inches(7.9), Inches(5.7))
veil.fill.solid()
veil.fill.fore_color.rgb = BG
_alpha(veil, 35)
veil.line.fill.background()
fade = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.6), 0, Inches(1.6), H)
fade.line.fill.background()
fade.fill.gradient()
fade.fill.gradient_angle = 0
st = fade.fill.gradient_stops
st[0].color.rgb, st[0].position = BG, 0
st[1].color.rgb, st[1].position = BG, 1.0
_g = st[1]._gs.find(qn("a:srgbClr"))
_g.append(_g.makeelement(qn("a:alpha"), {"val": "0"}))
text(s, Inches(0.7), Inches(1.3), Inches(6), Inches(0.4), "SCOPE · VIT CHENNAI", size=13,
     color=CYAN, bold=True)
text(s, Inches(0.7), Inches(1.8), Inches(6.2), Inches(1.0), [[("Medi", {"color": TEXT}), ("Twin", {"color": CYAN})]],
     size=66, bold=True)
text(s, Inches(0.7), Inches(3.0), Inches(5.2), Inches(1.4),
     "AI-Powered Hospital Digital Twin for Predictive Healthcare Operations and Decision Intelligence",
     size=22, color=TEXT, spacing=1.15)
x = Inches(0.7)
for lbl, c in (("Forecast", CYAN), ("Simulate", PINK), ("Optimise", AMBER), ("Recommend", GREEN)):
    x += pill(s, x, Inches(4.55), lbl, c) + Inches(0.12)
text(s, Inches(0.7), Inches(5.45), Inches(5.6), Inches(0.9),
     [[("Shresth Kumar Gupta", {"bold": True}), ("  ·  23BCE1578", {"color": MUTED})],
      [("Arnav Kamra", {"bold": True}), ("  ·  23BCE1706", {"color": MUTED})]], size=16, spacing=1.3)
text(s, Inches(0.7), Inches(6.55), Inches(5.6), Inches(0.3), "School of Computer Science and Engineering",
     size=12, color=GREY)

# ------------------------------------------------------------------ 2 problem
s = new_slide("Hospitals run on yesterday's numbers", "The problem", PINK)
bullets(s, Inches(0.6), Inches(1.9), Inches(6.3), [
    [("Reactive operations. ", {"bold": True}), "Bed crises, ED crowding and cancelled surgery are noticed "
                                               "only after they happen."],
    [("Siloed decisions. ", {"bold": True}), "Nursing rosters, theatre lists and bed management are planned "
                                             "separately, though they compete for the same beds."],
    [("No safe place to test. ", {"bold": True}), "A new policy (surge beds, diversion, deferring electives) "
                                                  "can only be tried on real patients."],
    [("Fragmented research. ", {"bold": True}), "Most work does forecasting, simulation or optimisation, "
                                                "rarely all three, and almost never verifies its advice."],
], size=17, bullet_color=PINK, gap=14)
stat(s, Inches(7.4), Inches(1.95), Inches(2.6), Inches(1.55), "100%", "Medical ward now", RED, "180 / 180 beds")
stat(s, Inches(10.2), Inches(1.95), Inches(2.6), Inches(1.55), "34 h", "ICU full next week", AMBER,
     "if nothing changes")
stat(s, Inches(7.4), Inches(3.7), Inches(2.6), Inches(1.55), "91", "Left without being seen", PINK, "per week, ED")
stat(s, Inches(10.2), Inches(3.7), Inches(2.6), Inches(1.55), "30", "Electives cancelled", VIOLET,
     "per week, no bed")
text(s, Inches(7.4), Inches(5.45), Inches(5.4), Inches(0.6),
     "Baseline 7-day projection of the twin from its live state (Wed 14 Oct 2026, 08:00).", size=11, color=GREY)

# ------------------------------------------------------------------ 3 objectives
s = new_slide("What MediTwin sets out to do", "Objectives", CYAN)
cards = [
    ("01", "Forecast", "Predict ED arrivals and ICU / ward / ED occupancy 14 days ahead, with calibrated "
                       "uncertainty and explanations.", CYAN),
    ("02", "Mirror", "Hold a live digital twin of every bed, queue and patient, and simulate the hospital "
                     "forward from that exact state.", PINK),
    ("03", "Optimise", "Turn forecasts into nurse rosters, theatre lists and patient transfers that "
                       "respect real constraints.", AMBER),
    ("04", "Recommend", "Propose actions and test each one in the twin first. Only advice that helps is "
                        "shown.", GREEN),
]
cw = Inches(2.9)
for i, (n, t, d, c) in enumerate(cards):
    x = Inches(0.6) + i * (cw + Inches(0.18))
    box(s, x, Inches(1.95), cw, Inches(3.4))
    text(s, x + Inches(0.3), Inches(2.2), cw, Inches(0.6), n, size=34, bold=True, color=c)
    text(s, x + Inches(0.3), Inches(2.95), cw - Inches(0.5), Inches(0.5), t, size=22, bold=True)
    text(s, x + Inches(0.3), Inches(3.55), cw - Inches(0.55), Inches(1.7), d, size=14, color=MUTED, spacing=1.2)
box(s, Inches(0.6), Inches(5.6), Inches(12.15), Inches(0.95), fill=RGBColor(0x0C, 0x22, 0x2B), line=CYAN)
text(s, Inches(0.9), Inches(5.6), Inches(11.6), Inches(0.95),
     [[("Novelty: ", {"bold": True, "color": CYAN}),
       ("all four layers in one loop. Every recommendation is re-simulated and reported with its KPI change and "
        "90% confidence interval before anyone sees it.", {})]], size=16, anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------------------ 4 architecture
s = new_slide("System architecture", "How it fits together", VIOLET)
y = Inches(2.2)
nw, nh = Inches(2.15), Inches(1.15)
xs = [Inches(0.6) + i * Inches(2.5) for i in range(5)]
row = [("Hospital data", "2 yrs · 165k encounters\nADT-style snapshot", GREY),
       ("Pre-processing", "hourly / daily series\ncalendar · holidays", MUTED),
       ("ML forecasting", "SARIMA · LightGBM\nensemble · SHAP", CYAN),
       ("Digital twin", "live state of every\nbed, queue, patient", PINK),
       ("Simulation", "SimPy discrete-event\nMonte Carlo · CRN", VIOLET)]
for x, (t, d, c) in zip(xs, row):
    node(s, x, y, nw, nh, t, d, c)
for a, b in zip(xs, xs[1:]):
    arrow(s, a + nw, y + nh // 2, b, y + nh // 2)
opt_x, rec_x, ui_x = Inches(5.6), Inches(9.2), Inches(9.2)
oy = Inches(4.2)
node(s, Inches(3.1), oy, Inches(3.2), nh, "Optimisation", "CP-SAT roster · CP-SAT electives\nLP redistribution",
     AMBER)
node(s, Inches(7.3), oy, Inches(3.2), nh, "AI recommendation engine", "detect → propose → simulate → score",
     GREEN)
arrow(s, xs[2] + nw // 2, y + nh, Inches(4.7), oy)
arrow(s, Inches(6.3), oy + nh // 2, Inches(7.3), oy + nh // 2)
arrow(s, xs[4] + nw // 2, y + nh, Inches(9.6), oy)
arrow(s, Inches(9.9), oy, xs[4] + Inches(0.6), y + nh, color=GREEN)
text(s, Inches(10.05), Inches(3.55), Inches(1.6), Inches(0.3), "verify", size=11, color=GREEN, bold=True)
node(s, Inches(11.0), oy, Inches(1.75), nh, "Interfaces", "Streamlit · FastAPI", CYAN)
arrow(s, Inches(10.5), oy + nh // 2, Inches(11.0), oy + nh // 2)
text(s, Inches(0.6), Inches(5.85), Inches(12), Inches(0.9),
     [[("One command builds the full twin in ~72 s: ", {"bold": True}),
       ("hospital-twin build → generate history → back-test 4 models × 5 targets → fit + explain → optimise → "
        "project 24 futures → verify recommendations.", {"color": MUTED})]], size=14)

# ------------------------------------------------------------------ 5 data
s = new_slide("A realistic synthetic hospital", "Data", AMBER)
stats = [("360", "beds", "ED 36 · ICU 24 · Med 180 · Surg 120", AMBER),
         ("164,594", "encounters", "730 days of history", CYAN),
         ("7", "theatres", "+ 16-chair discharge lounge", VIOLET),
         ("5", "ESI levels", "triage, LWBS patience, boarding", PINK)]
for i, (v, l, sub, c) in enumerate(stats):
    stat(s, Inches(0.6) + i * Inches(3.07), Inches(1.9), Inches(2.9), Inches(1.5), v, l, c, sub)
text(s, Inches(0.6), Inches(3.75), Inches(6), Inches(0.4), "Demand reflects Chennai", size=18, bold=True)
bullets(s, Inches(0.6), Inches(4.25), Inches(6), [
    "Weekly cycle and hour-of-day profile",
    "North-east monsoon fever season, summer heat, winter respiratory wave",
    "Indian public holidays and a Diwali trauma spike",
    "AR(1) day-to-day noise and random mass-casualty events",
], size=14, bullet_color=AMBER, gap=6)
text(s, Inches(6.9), Inches(3.75), Inches(6), Inches(0.4), "Why it is credible", size=18, bold=True)
bullets(s, Inches(6.9), Inches(4.25), Inches(5.9), [
    "History is produced by running demand through the simulator, so occupancy, boarding and cancellations "
    "respect real capacity limits",
    "Fictional hospital (Aarogya Multispecialty, Chennai): no patient data used",
    "Real HL7 / FHIR ADT feeds would replace the generator without other changes",
], size=14, bullet_color=AMBER, gap=6)

# ------------------------------------------------------------------ 6 forecasting
s = new_slide("Forecasting with calibrated uncertainty", "ML forecasting", CYAN)
picture(s, IMG / "forecasts.png", Inches(0.6), Inches(1.85), w=Inches(6.6))
bullets(s, Inches(7.55), Inches(1.9), Inches(5.2), [
    [("4 models: ", {"bold": True}), "seasonal naive, SARIMA with holiday regressors, recursive LightGBM, "
                                     "ensemble"],
    [("Rolling back-test: ", {"bold": True}), "10 origins × 14-day horizon picks a champion per target"],
    [("80% bands ", {"bold": True}), "calibrated on back-test residuals"],
    [("SHAP ", {"bold": True}), "explains the drivers of every forecast day"],
], size=14, gap=6)
table(s, Inches(7.55), Inches(4.35), Inches(5.2), [
    ["Target", "Champion", "MAPE", "Skill", "Cover"],
    ["ED arrivals", "LightGBM", "8.8%", "+23%", "79%"],
    ["ED census", "LightGBM", "10.6%", "+28%", "81%"],
    ["ICU census", "Ensemble", "13.2%", "+18%", "79%"],
    ["Medical ward", "Ensemble", "6.5%", "+19%", "82%"],
    ["Surgical ward", "LightGBM", "9.9%", "+24%", "81%"],
], col_w=[1.45, 1.1, 0.9, 0.85, 0.9], size=12, row_h=0.33, highlight_col=3)
text(s, Inches(7.55), Inches(6.45), Inches(5.2), Inches(0.3),
     "Skill = MAE reduction vs seasonal naive. Cover = actuals inside the 80% band.", size=10, color=GREY)

# ------------------------------------------------------------------ 7 twin
s = new_slide("The live digital twin", "Digital twin + simulation", PINK)
picture(s, IMG / "home.png", Inches(0.6), Inches(1.85), w=Inches(6.9))
bullets(s, Inches(8.25), Inches(1.9), Inches(4.6), [
    [("Snapshot ", {"bold": True}), "of every bed, queue and patient: unit, elapsed stay, acuity"],
    [("Warm start: ", {"bold": True}), "simulations begin from now, sampling remaining LOS from truncated "
                                       "log-normals"],
    [("SimPy model: ", {"bold": True}), "ESI triage, LWBS, doctor shifts, boarding and outliers, emergency and "
                                        "elective theatres, ICU exit-block, discharge window, lounge"],
    [("Monte Carlo: ", {"bold": True}), "24 futures, demand drawn from the forecast band, common random "
                                        "numbers for fair comparisons"],
    [("Animated floor plan ", {"bold": True}), "plays the next 72 h bed by bed"],
], size=14, bullet_color=PINK, gap=7)

# ------------------------------------------------------------------ 8 scenario lab
s = new_slide("Scenario Lab: stress-test before it happens", "What-if simulation", PINK)
picture(s, IMG / "lab_twin.png", Inches(0.6), Inches(1.85), w=Inches(6.15))
picture(s, IMG / "lab_traj.png", Inches(6.95), Inches(1.85), w=Inches(5.8))
text(s, Inches(0.6), Inches(5.75), Inches(12.2), Inches(1),
     [[("Mass casualty preset: ", {"bold": True, "color": PINK}),
       ("35 high-acuity patients arrive 2 hours into day 2. Door-to-doctor rises to 37 min and ICU is full for 70 h. "
        "Other presets: monsoon fever surge, ICU bed closure, staff shortage, festival week. One click generates "
        "an AI response plan for the scenario.", {"color": MUTED})]], size=14, spacing=1.2)

# ------------------------------------------------------------------ 9 optimisation
s = new_slide("Optimisation: from forecast to plan", "Operations research", AMBER)
picture(s, IMG / "optimizer.png", Inches(0.6), Inches(1.85), w=Inches(6.6))
cols = [("Nurse roster", "OR-Tools CP-SAT", "≤ 1 shift / day, no night → day, contracted hours, night cap, "
                                            "float pool, agency", "0", "uncovered shifts", "template: 130"),
        ("Elective list", "OR-Tools CP-SAT", "theatre hours + post-op ICU and ward beds against forecast "
                                             "emergency load", "0", "bed-overflow days", "FIFO: 11"),
        ("Redistribution", "SciPy LP (HiGHS)", "moves excess patients to outlier beds or partner hospitals at "
                                               "minimum cost", "LP", "optimal transfers", "")]
for i, (t, eng, d, v, l, cmp) in enumerate(cols):
    y = Inches(1.85) + i * Inches(1.6)
    box(s, Inches(7.55), y, Inches(5.2), Inches(1.45))
    text(s, Inches(7.8), y + Inches(0.15), Inches(3), Inches(0.35), t, size=16, bold=True, color=AMBER)
    text(s, Inches(7.8), y + Inches(0.48), Inches(3.1), Inches(0.3), eng, size=11, color=CYAN, bold=True)
    text(s, Inches(7.8), y + Inches(0.75), Inches(3.15), Inches(0.7), d, size=10.5, color=MUTED)
    text(s, Inches(11.0), y + Inches(0.15), Inches(1.6), Inches(0.6), v, size=30, bold=True, color=GREEN,
         align=PP_ALIGN.CENTER)
    text(s, Inches(11.0), y + Inches(0.75), Inches(1.6), Inches(0.3), l, size=10, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, Inches(11.0), y + Inches(1.0), Inches(1.6), Inches(0.3), cmp, size=10, color=RED, align=PP_ALIGN.CENTER)

# ------------------------------------------------------------------ 10 recommender
s = new_slide("AI recommendations, verified by simulation", "Decision intelligence", GREEN)
steps = [("Detect", "risks in forecasts and twin projections", CYAN),
         ("Propose", "playbook actions with owner and urgency", VIOLET),
         ("Simulate", "re-run the twin with each action (CRN)", PINK),
         ("Score", "weighted KPI benefit with 90% CI", AMBER),
         ("Bundle", "combine winners and re-verify as one plan", GREEN)]
for i, (t, d, c) in enumerate(steps):
    x = Inches(0.6) + i * Inches(2.47)
    node(s, x, Inches(1.9), Inches(2.15), Inches(1.0), t, d, c)
    if i:
        arrow(s, x - Inches(0.32), Inches(2.4), x, Inches(2.4))
recs = [("Accept", "Integrated response plan", "ICU full 34.5 → 1.8 h", "185", GREEN),
        ("Accept", "Open 4 temporary ICU beds", "ICU full −75%", "85", GREEN),
        ("Accept", "Discharge-before-noon drive", "ICU full −30%", "85", GREEN),
        ("Accept", "Bed-aware elective schedule", "electives +40%", "73", GREEN),
        ("Accept", "Optimised nurse roster", "0 uncovered shifts", "18", GREEN),
        ("Reject", "Extra ED physician", "cost > benefit", "−17", RED),
        ("Reject", "Defer 30% of electives", "lost surgery outweighs gain", "−51", RED)]
table(s, Inches(0.6), Inches(3.25), Inches(7.2),
      [["Verdict", "Action", "Simulated effect", "Score"]] + [[a, b, c, d] for a, b, c, d, _ in recs],
      col_w=[0.95, 2.95, 2.45, 0.85], size=12, row_h=0.4, left_cols=(1,), colors={0: {r[0]: r[4] for r in recs}})
picture(s, IMG / "recommendations.png", Inches(8.05), Inches(3.25), w=Inches(4.7))

# ------------------------------------------------------------------ 11 ablation chart
s = new_slide("Results: AI plan vs reactive operations", "Ablation", GREEN)
cd = CategoryChartData()
cd.categories = ["Baseline", "Monsoon surge", "Mass casualty", "ICU closure"]
cd.add_series("Reactive", (34.5, 92.4, 70.2, 125.0))
cd.add_series("AI plan", (1.8, 26.4, 7.9, 14.5))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.8), Inches(7.6), Inches(4.9), cd)
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.legend.font.color.rgb = TEXT
ch.legend.font.size = Pt(12)
ch.legend.font.name = FONT
ch.has_title = True
ch.chart_title.text_frame.text = "Hours with ICU full (next 7 days, mean of 24 futures)"
tp = ch.chart_title.text_frame.paragraphs[0]
tp.runs[0].font.size, tp.runs[0].font.color.rgb, tp.runs[0].font.name = Pt(13), MUTED, FONT
plot = ch.plots[0]
plot.gap_width, plot.overlap = 60, -10
plot.has_data_labels = True
dl = plot.data_labels
dl.font.size, dl.font.color.rgb, dl.font.bold, dl.font.name = Pt(12), TEXT, True, FONT
dl.number_format, dl.number_format_is_linked = "0", False
for ser, c in zip(plot.series, (GREY, GREEN)):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = c
for ax in (ch.category_axis, ch.value_axis):
    ax.tick_labels.font.color.rgb = MUTED
    ax.tick_labels.font.size = Pt(12)
    ax.tick_labels.font.name = FONT
    ax.format.line.color.rgb = LINE
ch.value_axis.major_gridlines.format.line.color.rgb = LINE
ch.value_axis.has_major_gridlines = True
table(s, Inches(8.55), Inches(1.85), Inches(4.2), [
    ["Under stress", "Reactive", "AI plan"],
    ["Door-to-doctor (MCI)", "37 min", "19 min"],
    ["ED boarding (ICU closure)", "2.3 h", "0.4 h"],
    ["LWBS (monsoon)", "606", "245"],
    ["LWBS (ICU closure)", "172", "24"],
    ["Electives cancelled (MCI)", "37", "21"],
], col_w=[2.2, 1.0, 1.0], size=12, row_h=0.42, highlight_col=2)
box(s, Inches(8.55), Inches(4.75), Inches(4.2), Inches(1.9), fill=RGBColor(0x0B, 0x24, 0x1E), line=GREEN)
text(s, Inches(8.8), Inches(4.9), Inches(3.8), Inches(1.6),
     [[("3.5–19×", {"size": 36, "bold": True, "color": GREEN})],
      [("fewer ICU-full hours across every stress scenario, with shorter ED waits and fewer patients "
        "leaving unseen.", {"color": TEXT})]], size=13, spacing=1.15)

# ------------------------------------------------------------------ 12 validation
s = new_slide("Can we trust the twin?", "Validation", VIOLET)
text(s, Inches(0.6), Inches(1.85), Inches(6), Inches(0.4), "Twin replay: last 8 weeks of observed demand",
     size=16, bold=True)
table(s, Inches(0.6), Inches(2.35), Inches(6.2), [
    ["Series", "Observed", "Twin", "MAPE"],
    ["ED arrivals / h", "8.92", "8.95", "6.1%"],
    ["ED census", "22.9", "23.2", "8.2%"],
    ["ICU census", "19.5", "20.3", "14.1%"],
    ["Medical ward", "153.3", "156.0", "5.0%"],
    ["Surgical ward", "89.1", "94.7", "10.2%"],
], col_w=[2.3, 1.3, 1.3, 1.3], size=13, header_color=VIOLET, row_h=0.4, highlight_col=3)
text(s, Inches(0.6), Inches(4.95), Inches(6.2), Inches(0.4), "Optimisers vs simple rules", size=16, bold=True)
table(s, Inches(0.6), Inches(5.4), Inches(6.2), [
    ["Metric", "Rule-based", "Optimised"],
    ["Uncovered nurse shifts (7 d)", "130", "0"],
    ["Electives: bed-overflow days", "11", "0"],
    ["Electives: P1 cases scheduled", "77%", "94%"],
], col_w=[3.2, 1.5, 1.5], size=13, header_color=VIOLET, row_h=0.36, highlight_col=2)
picture(s, IMG / "model_lab.png", Inches(7.2), Inches(1.85), w=Inches(5.55))
text(s, Inches(7.2), Inches(6.0), Inches(5.55), Inches(0.7),
     "Model Lab page: architecture, back-test detail, twin validation, optimisation benchmarks and ablation, "
     "regenerated by hospital-twin evaluate.", size=11, color=GREY)

# ------------------------------------------------------------------ 13 command centre
s = new_slide("The command centre", "Interfaces", CYAN)
picture(s, IMG / "patient_flow.png", Inches(0.6), Inches(1.85), w=Inches(6.0))
picture(s, IMG / "lab_sankey.png", Inches(6.8), Inches(1.85), w=Inches(5.95))
pages = ["Command Center", "Forecasts", "Scenario Lab", "Optimizer", "AI Recommendations", "Patient Flow",
         "Model Lab"]
x = Inches(0.6)
for p_ in pages:
    x += pill(s, x, Inches(6.35), p_, CYAN, w=Inches(0.105 * len(p_) + 0.45)) + Inches(0.1)
text(s, Inches(6.8), Inches(5.0), Inches(5.95), Inches(1.1),
     [[("REST API (FastAPI): ", {"bold": True, "color": CYAN}),
       ("GET /state · /forecast/{target} · /roster · /electives\nPOST /simulate · /recommendations", {})]],
     size=13, spacing=1.25)

# ------------------------------------------------------------------ 14 stack
s = new_slide("Technology stack", "Implementation", VIOLET)
stack = [("Language & build", "Python 3.12 · uv · pytest (17 tests)", MUTED),
         ("Forecasting", "statsmodels SARIMAX · LightGBM · SHAP · scikit-learn", CYAN),
         ("Simulation", "SimPy discrete-event · joblib parallel Monte Carlo", PINK),
         ("Optimisation", "Google OR-Tools CP-SAT · SciPy linprog (HiGHS)", AMBER),
         ("Decision layer", "simulation-verified playbook · optional Claude-written executive brief", GREEN),
         ("Interfaces", "Streamlit · Plotly (animated floor plan, Sankey, maps) · FastAPI", VIOLET)]
for i, (k, v, c) in enumerate(stack):
    y = Inches(1.9) + i * Inches(0.75)
    box(s, Inches(0.6), y, Inches(12.15), Inches(0.62))
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.85), y + Inches(0.22), Inches(0.18), Inches(0.18))
    dot.fill.solid()
    dot.fill.fore_color.rgb = c
    dot.line.fill.background()
    text(s, Inches(1.25), y, Inches(3), Inches(0.62), k, size=16, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(4.3), y, Inches(8.3), Inches(0.62), v, size=15, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(0.6), Inches(6.5), Inches(12), Inches(0.4),
     "uv run hospital-twin build  ·  evaluate  ·  app  ·  api", size=13, color=CYAN, bold=True)

# ------------------------------------------------------------------ 15 limitations
s = new_slide("Limitations and future work", "Next steps", AMBER)
text(s, Inches(0.6), Inches(1.85), Inches(5.9), Inches(0.4), "Limitations", size=18, bold=True, color=AMBER)
bullets(s, Inches(0.6), Inches(2.35), Inches(5.9), [
    "Patient data is synthetic. Volumes are plausible for an Indian tertiary hospital, but it is not a real "
    "record.",
    "Twin validation is partly self-consistent, because history comes from the same simulator.",
    "Nurse staffing is optimised but does not slow care inside the simulator (ED doctors do).",
    "Partner-hospital capacity is static and fictional.",
], size=15, bullet_color=AMBER, gap=10)
text(s, Inches(6.9), Inches(1.85), Inches(5.9), Inches(0.4), "Future work", size=18, bold=True, color=GREEN)
bullets(s, Inches(6.9), Inches(2.35), Inches(5.9), [
    "Connect live HL7 / FHIR ADT feeds and validate against a partner hospital.",
    "Model nurse workload in the simulator, so roster gaps affect patient flow.",
    "Reinforcement learning for sequential bed-management policies.",
    "Live regional bed availability for inter-hospital transfers.",
], size=15, bullet_color=GREEN, gap=10)

# ------------------------------------------------------------------ 16 thanks
s = new_slide(accent=GREEN)
text(s, 0, Inches(2.2), W, Inches(1.2), [[("Thank ", {"color": TEXT}), ("you", {"color": CYAN})]], size=72,
     bold=True, align=PP_ALIGN.CENTER)
text(s, 0, Inches(3.65), W, Inches(0.5), "Forecast  →  Simulate  →  Optimise  →  Recommend", size=20, color=MUTED,
     align=PP_ALIGN.CENTER)
text(s, 0, Inches(4.6), W, Inches(0.4), "github.com/Arnav717/hospital-digital-twin", size=16, color=CYAN,
     bold=True, align=PP_ALIGN.CENTER)
text(s, 0, Inches(5.5), W, Inches(0.8),
     ["Shresth Kumar Gupta (23BCE1578)  ·  Arnav Kamra (23BCE1706)", "SCOPE, VIT Chennai"], size=14, color=GREY,
     align=PP_ALIGN.CENTER, spacing=1.3)

assert _n == TOTAL, _n
out = HERE / "MediTwin.pptx"
prs.save(out)
print("wrote", out)
